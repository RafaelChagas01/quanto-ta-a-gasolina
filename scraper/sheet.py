import io
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import load_workbook

from .text import key

PRODUCTS = {
    "GASOLINA COMUM": "gasolina",
    "GASOLINA ADITIVADA": "gasolina_aditivada",
    "ETANOL HIDRATADO": "etanol",
    "OLEO DIESEL": "diesel",
    "OLEO DIESEL S10": "diesel_s10",
    "GNV": "gnv",
    "GLP": "glp",
}

REQUIRED = {
    "DATA FINAL": "end",
    "ESTADO": "state",
    "MUNICIPIO": "city",
    "PRODUTO": "product",
    "NUMERO DE POSTOS PESQUISADOS": "stations",
    "PRECO MEDIO REVENDA": "avg",
    "PRECO MINIMO REVENDA": "min",
    "PRECO MAXIMO REVENDA": "max",
}

MAX_PRICE = 1000.0


@dataclass(frozen=True)
class PriceRow:
    end: date
    state: str
    city: str
    product: str
    stations: int
    avg: float
    min: float
    max: float


class SheetFormatError(Exception):
    pass


def _to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _to_float(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        try:
            return float(text)
        except ValueError:
            return None
    return None


def _clean_row(raw: dict) -> PriceRow | None:
    product = PRODUCTS.get(key(str(raw["product"] or "")))
    end = _to_date(raw["end"])
    state = key(str(raw["state"] or ""))
    city = key(str(raw["city"] or ""))
    stations = _to_float(raw["stations"])
    avg, lo, hi = (_to_float(raw[k]) for k in ("avg", "min", "max"))

    if not (product and end and state and city):
        return None
    if stations is None or stations < 1:
        return None
    if None in (avg, lo, hi):
        return None
    if not (0 < lo <= avg <= hi < MAX_PRICE):
        return None

    return PriceRow(end, state, city, product, int(stations), round(avg, 2), round(lo, 2), round(hi, 2))


def read_city_prices(content: bytes) -> tuple[list[PriceRow], int]:
    """Le a aba MUNICIPIOS. Retorna as linhas validas e quantas foram descartadas."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = next((ws for ws in wb.worksheets if key(ws.title) == "MUNICIPIOS"), None)
    if sheet is None:
        raise SheetFormatError("aba MUNICIPIOS nao encontrada")

    columns: dict[str, int] | None = None
    rows: list[PriceRow] = []
    skipped = 0

    for values in sheet.iter_rows(values_only=True):
        if columns is None:
            labels = [key(str(v)) if v is not None else "" for v in values]
            if "PRODUTO" in labels and "MUNICIPIO" in labels:
                missing = [name for name in REQUIRED if name not in labels]
                if missing:
                    raise SheetFormatError(f"colunas ausentes: {', '.join(missing)}")
                columns = {field: labels.index(name) for name, field in REQUIRED.items()}
            continue

        if not any(v is not None for v in values):
            continue

        raw = {field: values[i] if i < len(values) else None for field, i in columns.items()}
        row = _clean_row(raw)
        if row is None:
            skipped += 1
        else:
            rows.append(row)

    wb.close()
    if columns is None:
        raise SheetFormatError("cabecalho nao encontrado")
    return rows, skipped
